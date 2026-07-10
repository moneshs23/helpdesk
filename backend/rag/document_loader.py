"""Text extraction from supported document formats.

Each loader returns a list of ``PageText`` items carrying page/section context so
that downstream chunks can be cited precisely (filename + page + section).

Row/record-oriented formats (Excel, CSV, JSON, XML) are paginated into blocks
of rows so page numbers stay meaningful for citation and progress display.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path

import chardet

from backend.config import settings
from backend.models.enums import DocumentType
from backend.utils.logging import logger

# Rows/records grouped into one logical "page" for tabular formats.
ROWS_PER_PAGE = 50


@dataclass
class PageText:
    page: int
    section: str
    text: str
    # True for tabular/record formats (Excel, CSV, JSON, XML) where each line
    # of `text` is one independent row/record. The chunker uses this to keep
    # every row as its own chunk instead of packing several rows together.
    row_oriented: bool = False


# ---------------- individual extractors ----------------
def _read_bytes(path: Path) -> bytes:
    return path.read_bytes()


def _decode(data: bytes) -> str:
    guess = chardet.detect(data)
    encoding = guess.get("encoding") or "utf-8"
    try:
        return data.decode(encoding, errors="replace")
    except (LookupError, UnicodeDecodeError):
        return data.decode("utf-8", errors="replace")


def _lines_to_pages(
    lines: list[str], section: str, start_page: int = 1, unit: str = "rows"
) -> list[PageText]:
    """Group flat lines into pages of ROWS_PER_PAGE for citable page numbers."""
    pages: list[PageText] = []
    for i in range(0, len(lines), ROWS_PER_PAGE):
        block = lines[i : i + ROWS_PER_PAGE]
        label = f"{section} ({unit} {i + 1}-{i + len(block)})" if section else ""
        pages.append(
            PageText(
                page=start_page + i // ROWS_PER_PAGE,
                section=label,
                text="\n".join(block),
                row_oriented=True,
            )
        )
    return pages


def load_txt(path: Path) -> list[PageText]:
    return [PageText(page=1, section="", text=_decode(_read_bytes(path)))]


def load_md(path: Path) -> list[PageText]:
    text = _decode(_read_bytes(path))
    pages: list[PageText] = []
    current_section = ""
    buffer: list[str] = []
    page_no = 1
    for line in text.splitlines():
        if line.startswith("#"):
            if buffer:
                pages.append(
                    PageText(page_no, current_section, "\n".join(buffer).strip())
                )
                buffer = []
                page_no += 1
            current_section = line.lstrip("# ").strip()
        buffer.append(line)
    if buffer:
        pages.append(PageText(page_no, current_section, "\n".join(buffer).strip()))
    return pages or [PageText(1, "", text)]


def load_pdf(path: Path) -> list[PageText]:
    from pypdf import PdfReader  # noqa: PLC0415

    reader = PdfReader(str(path))
    pages: list[PageText] = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text and settings.enable_ocr:
            text = _ocr_pdf_page(path, i - 1)
        if text:
            pages.append(PageText(page=i, section="", text=text))
    return pages


def _ocr_pdf_page(path: Path, page_index: int) -> str:
    """Best-effort OCR for scanned PDF pages (requires tesseract + pdf2image)."""
    try:
        import pytesseract  # noqa: PLC0415
        from pdf2image import convert_from_path  # noqa: PLC0415

        images = convert_from_path(
            str(path), first_page=page_index + 1, last_page=page_index + 1, dpi=200
        )
        if images:
            return pytesseract.image_to_string(images[0]).strip()
    except Exception as exc:  # noqa: BLE001
        logger.warning(f"OCR failed for {path.name} page {page_index + 1}: {exc}")
    return ""


def load_docx(path: Path) -> list[PageText]:
    from docx import Document as DocxDocument  # noqa: PLC0415

    try:
        doc = DocxDocument(str(path))
    except Exception as exc:
        if path.suffix.lower() == ".doc":
            raise ValueError(
                "Legacy .doc files are not supported. Please re-save the file "
                "as .docx and upload again."
            ) from exc
        raise
    pages: list[PageText] = []
    current_section = ""
    buffer: list[str] = []
    page_no = 1
    for para in doc.paragraphs:
        style = (para.style.name or "").lower() if para.style else ""
        if style.startswith("heading"):
            if buffer:
                pages.append(
                    PageText(page_no, current_section, "\n".join(buffer).strip())
                )
                buffer = []
                page_no += 1
            current_section = para.text.strip()
        if para.text.strip():
            buffer.append(para.text.strip())
    # Tables
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                buffer.append(" | ".join(cells))
    if buffer:
        pages.append(PageText(page_no, current_section, "\n".join(buffer).strip()))
    return pages or [PageText(1, "", "")]


def load_csv(path: Path) -> list[PageText]:
    import pandas as pd  # noqa: PLC0415

    sep = "\t" if path.suffix.lower() == ".tsv" else ","
    df = pd.read_csv(path, dtype=str, keep_default_na=False, sep=sep)
    return _lines_to_pages(_dataframe_to_lines(df), section="CSV data")


def load_xlsx(path: Path) -> list[PageText]:
    import pandas as pd  # noqa: PLC0415

    sheets = pd.read_excel(path, sheet_name=None, dtype=str)
    pages: list[PageText] = []
    next_page = 1
    for name, df in sheets.items():
        df = df.fillna("")
        lines = _dataframe_to_lines(df)
        if not lines:
            continue
        sheet_pages = _lines_to_pages(
            lines, section=f"Sheet: {name}", start_page=next_page
        )
        pages.extend(sheet_pages)
        next_page = pages[-1].page + 1
    return pages or [PageText(1, "", "")]


# Ticket-history narrative like "Open (31.12.2025) : 15:37:49 : <question>
# Assigned (...) : ... Response (...) : ... Closed (05.01.2026) : <answer>"
# packed into a single cell (e.g. a helpdesk export's "Remarks" column).
_TICKET_EVENT_RE = re.compile(
    r"(Open|Assigned|Response|Re\s*Open|Reopen|Partially\s*Closed|Closed)"
    r"\s*\([^)]{0,40}\)\s*:\s*(?:[0-9]{1,2}:[0-9]{2}:[0-9]{2}\s*:\s*)?",
    re.IGNORECASE,
)


def _split_ticket_narrative(text: str) -> dict[str, list[str]]:
    """Split a ticket-history cell into its named event segments (open,
    assigned, response, closed, ...), discarding nothing but re-keying by
    event so the caller can pick the first "open" and the last "closed".
    Returns {} if the text has no recognizable Open/Closed markers.
    """
    matches = list(_TICKET_EVENT_RE.finditer(text))
    events: dict[str, list[str]] = {}
    for i, m in enumerate(matches):
        label = re.sub(r"\s+", "", m.group(1)).lower()
        start = m.end()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        content = text[start:end].strip()
        if content == "-":
            content = ""
        if content:
            events.setdefault(label, []).append(content)
    return events


def _find_narrative_column(df) -> str | None:
    """Detect the one column (if any) holding an open/closed ticket
    narrative. Returns None for ordinary tabular data with no such column."""
    best_col, best_hits = None, 0
    for col in df.columns:
        hits = 0
        for value in df[col]:
            events = _split_ticket_narrative(str(value))
            if events.get("open") and events.get("closed"):
                hits += 1
        if hits > best_hits:
            best_col, best_hits = col, hits
    return best_col


def _dataframe_to_lines(df) -> list[str]:
    """Render each row as 'header: value; ...', skipping pandas' auto headers.

    If a column holds an open/closed ticket narrative (e.g. a support-ticket
    export's "Remarks" column), only that column's opened/closed text is kept
    — plus the VIN, which is a unique per-vehicle identifier that makes
    exact-match retrieval deterministic (a query quoting the VIN can only hit
    its own ticket). Every other column (ticket IDs, dates, part-number
    metadata, ...) is dropped as noise that dilutes the row's embedding.
    """
    narrative_col = _find_narrative_column(df)
    if narrative_col is not None:
        vin_col = next(
            (c for c in df.columns if str(c).strip().lower() == "vin"), None
        )
        lines: list[str] = []
        for _, row in df.iterrows():
            events = _split_ticket_narrative(str(row[narrative_col]))
            opened = events["open"][0] if events.get("open") else ""
            closed = events["closed"][-1] if events.get("closed") else ""
            parts = []
            if vin_col is not None:
                vin = str(row[vin_col]).strip()
                if vin and vin not in ("-", "nan"):
                    parts.append(f"VIN: {vin}")
            if opened:
                parts.append(f"Opened: {opened}")
            if closed:
                parts.append(f"Closed: {closed}")
            if opened or closed:
                lines.append("; ".join(parts))
        return lines

    headers = list(df.columns)
    lines: list[str] = []
    for _, row in df.iterrows():
        parts = []
        for h in headers:
            value = str(row[h]).strip()
            if not value or value.lower() == "nan":
                continue
            label = str(h).strip()
            if label.startswith("Unnamed:"):
                parts.append(value)
            else:
                parts.append(f"{label}: {value}")
        if parts:
            lines.append("; ".join(parts))
    return lines


def load_pptx(path: Path) -> list[PageText]:
    from pptx import Presentation  # noqa: PLC0415

    try:
        prs = Presentation(str(path))
    except Exception as exc:
        if path.suffix.lower() == ".ppt":
            raise ValueError(
                "Legacy .ppt files are not supported. Please re-save the file "
                "as .pptx and upload again."
            ) from exc
        raise
    pages: list[PageText] = []
    for i, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
        title = ""
        try:
            if slide.shapes.title is not None:
                title = slide.shapes.title.text.strip()
        except Exception:  # noqa: BLE001 - layouts without a title placeholder
            title = ""
        if parts:
            pages.append(PageText(page=i, section=title, text="\n".join(parts)))
    return pages


class _HTMLTextExtractor(HTMLParser):
    """Strip tags, skip script/style, and split sections on h1-h3 headings."""

    _SKIP = {"script", "style", "noscript", "template", "svg"}
    _HEADINGS = {"h1", "h2", "h3"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._skip_depth = 0
        self._in_heading = False
        self._heading_buf: list[str] = []
        self.blocks: list[tuple[str, list[str]]] = [("", [])]

    def handle_starttag(self, tag, attrs) -> None:  # noqa: ANN001
        if tag in self._SKIP:
            self._skip_depth += 1
        elif tag in self._HEADINGS:
            self._in_heading = True
            self._heading_buf = []

    def handle_endtag(self, tag) -> None:  # noqa: ANN001
        if tag in self._SKIP:
            self._skip_depth = max(0, self._skip_depth - 1)
        elif tag in self._HEADINGS and self._in_heading:
            self._in_heading = False
            section = " ".join(self._heading_buf).strip()
            self.blocks.append((section, [section] if section else []))

    def handle_data(self, data) -> None:  # noqa: ANN001
        if self._skip_depth:
            return
        text = data.strip()
        if not text:
            return
        if self._in_heading:
            self._heading_buf.append(text)
        else:
            self.blocks[-1][1].append(text)


def load_html(path: Path) -> list[PageText]:
    parser = _HTMLTextExtractor()
    parser.feed(_decode(_read_bytes(path)))
    pages: list[PageText] = []
    page_no = 1
    for section, lines in parser.blocks:
        text = "\n".join(lines).strip()
        if text:
            pages.append(PageText(page=page_no, section=section, text=text))
            page_no += 1
    return pages


def load_json(path: Path) -> list[PageText]:
    import json  # noqa: PLC0415

    data = json.loads(_decode(_read_bytes(path)))
    lines: list[str] = []

    def walk(node, prefix: str) -> None:  # noqa: ANN001
        if isinstance(node, dict):
            for key, value in node.items():
                walk(value, f"{prefix}.{key}" if prefix else str(key))
        elif isinstance(node, list):
            for i, value in enumerate(node):
                walk(value, f"{prefix}[{i}]")
        else:
            value = str(node).strip()
            if value:
                lines.append(f"{prefix}: {value}" if prefix else value)

    walk(data, "")
    return _lines_to_pages(lines, section="JSON data", unit="entries")


def load_xml(path: Path) -> list[PageText]:
    import xml.etree.ElementTree as ET  # noqa: PLC0415

    root = ET.fromstring(_decode(_read_bytes(path)))
    lines: list[str] = []

    def walk(el, prefix: str) -> None:  # noqa: ANN001
        tag = el.tag.split("}")[-1] if isinstance(el.tag, str) else ""
        node_path = f"{prefix}/{tag}" if prefix else tag
        for key, value in el.attrib.items():
            if str(value).strip():
                lines.append(f"{node_path}@{key}: {value}")
        text = (el.text or "").strip()
        if text:
            lines.append(f"{node_path}: {text}")
        for child in el:
            walk(child, node_path)

    walk(root, "")
    return _lines_to_pages(lines, section="XML data", unit="entries")


def load_rtf(path: Path) -> list[PageText]:
    from striprtf.striprtf import rtf_to_text  # noqa: PLC0415

    text = rtf_to_text(_decode(_read_bytes(path))).strip()
    return [PageText(page=1, section="", text=text)]


_LOADERS = {
    DocumentType.TXT: load_txt,
    DocumentType.MD: load_md,
    DocumentType.PDF: load_pdf,
    DocumentType.DOCX: load_docx,
    DocumentType.CSV: load_csv,
    DocumentType.XLSX: load_xlsx,
    DocumentType.PPTX: load_pptx,
    DocumentType.HTML: load_html,
    DocumentType.JSON: load_json,
    DocumentType.XML: load_xml,
    DocumentType.RTF: load_rtf,
}


def load_document(path: Path, doc_type: DocumentType) -> list[PageText]:
    loader = _LOADERS.get(doc_type)
    if loader is None:
        raise ValueError(f"No loader for document type: {doc_type}")
    pages = loader(path)
    return [p for p in pages if p.text and p.text.strip()]
